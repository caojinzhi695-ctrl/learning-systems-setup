"""Locally persisted review papers based only on the selected saved AI note."""
import json
from io import BytesIO

TYPES = {'recall': '简答题', 'fill': '填空题', 'single': '单选题', 'multiple': '多选题', 'correction': '判断并改错', 'explain': '对比与解释', 'application': '情境应用'}


def validate_paper(raw, finish):
    if finish not in {'stop', 'end_turn'}:
        raise ValueError('回答未完整结束，已保留草稿。')
    try:
        paper = json.loads(raw)
    except ValueError:
        raise ValueError('试卷格式不完整，已保留草稿。') from None
    if not isinstance(paper, dict):
        raise ValueError('试卷格式无效。')
    if paper.get('error'):
        raise ValueError('材料不足，无法生成试卷。')
    questions = paper.get('questions')
    if not isinstance(questions, list) or not 1 <= len(questions) <= 22:
        raise ValueError('题量必须为1—22题，已保留草稿。')
    for key in ('title', 'scope'):
        if not isinstance(paper.get(key), str) or not paper[key].strip():
            raise ValueError('试卷标题或范围缺失。')
    if not isinstance(paper.get('recap'), list) or not all(isinstance(x, str) for x in paper['recap']):
        raise ValueError('复习速记格式无效。')
    seen = set()
    for i, q in enumerate(questions, 1):
        if not isinstance(q, dict) or q.get('id') != i or q.get('type') not in TYPES:
            raise ValueError('题号或题型无效。')
        for key in ('prompt', 'answer', 'explanation', 'knowledgePoint'):
            if not isinstance(q.get(key), str) or not q[key].strip():
                raise ValueError(f'第{i}题题干或答案解析缺失。')
        prompt = ''.join(q['prompt'].split())
        if prompt in seen:
            raise ValueError('存在重复题目，已保留草稿。')
        seen.add(prompt)
        options = q.get('options', [])
        correct = q.get('correctOptions', [])
        if not isinstance(options, list) or not isinstance(correct, list):
            raise ValueError('选项格式无效。')
        if q['type'] in {'single', 'multiple'}:
            if not 2 <= len(options) <= 6 or any(not isinstance(o, dict) or o.get('key') != chr(65+j) or not isinstance(o.get('text'), str) or not o['text'].strip() for j, o in enumerate(options)):
                raise ValueError('选择题选项不完整。')
            if not all(isinstance(k, str) for k in correct) or len(set(correct)) != len(correct) or any(k not in [o['key'] for o in options] for k in correct) or (q['type'] == 'single' and len(correct) != 1) or (q['type'] == 'multiple' and len(correct) < 2):
                raise ValueError('选择题答案无效。')
        elif options or correct:
            raise ValueError('非选择题选项应为空。')
        if q['type'] == 'correction' and (not isinstance(q.get('verdict'), bool) or (q['verdict'] is False and not q.get('correction'))):
            raise ValueError('判断改错题答案不完整。')
    return paper


def answer(q):
    parts = []
    if q.get('correctOptions'):
        parts.append('正确选项：' + '、'.join(q['correctOptions']))
    if isinstance(q.get('verdict'), bool):
        parts.append('判断：' + ('正确' if q['verdict'] else '错误'))
    parts.append(q['answer'])
    if q.get('correction'):
        parts.append('正确表述：' + q['correction'])
    return '\n'.join(parts)


def export_word(paper):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(11)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    doc.add_paragraph(paper['title'] + ' 复习练习卷', 'Title')
    doc.add_paragraph('复习范围：' + paper['scope'])
    doc.add_paragraph(f"题量：{len(paper['questions'])}。先完成题目，再查看答案。")
    doc.add_heading('第一部分 题目', 1)
    for q in paper['questions']:
        doc.add_paragraph(f"{q['id']}. 【{TYPES[q['type']]}】{q['prompt']}")
        for o in q.get('options', []):
            doc.add_paragraph(o['key'] + '. ' + o['text'])
        if not q.get('options'):
            doc.add_paragraph('答：________________________________________________')
    doc.add_page_break()
    doc.add_heading('第二部分 参考答案与解析', 1)
    for q in paper['questions']:
        doc.add_heading(f"第{q['id']}题", 2)
        doc.add_paragraph('参考答案：' + answer(q))
        doc.add_paragraph('简要解析：' + q['explanation'])
        doc.add_paragraph('知识点：' + q['knowledgePoint'])
    doc.add_heading('第三部分 复习速记', 1)
    for line in paper['recap']:
        doc.add_paragraph(line, 'List Bullet')
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def interactive_note(paper, note_id=None):
    """Adapt a saved paper to the user's NOTE schema without another model call."""
    import hashlib
    import math
    digest = hashlib.sha256(json.dumps(paper, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:24]
    questions = []
    points = {}
    for q in paper['questions']:
        kind = q['type']
        is_choice = kind in ('single', 'multiple') and bool(q.get('options'))
        options = q.get('options', [])
        keys = q.get('correctOptions', [])
        correct = [i for i, option in enumerate(options) if option['key'] in keys]
        if is_choice and (not correct or (kind == 'single' and len(correct) != 1)):
            raise ValueError('试卷选项与参考答案不一致，请先核对原试卷。')
        entry = {'type': kind if is_choice else ('text' if kind == 'fill' else 'open'),
                 'label': TYPES.get(kind, kind), 'title': q['prompt'], 'answer': answer(q),
                 'why': q['explanation'], 'point': q['knowledgePoint']}
        if is_choice:
            entry.update(options=[o['text'] for o in options], correct=correct)
        # Free text is self-assessed; no invented exact-match or semantic grading.
        questions.append(entry)
        point = q['knowledgePoint'].strip()
        if point:
            points.setdefault(point, [])
            if q['explanation'] not in points[point]: points[point].append(q['explanation'])
    summary = [{'label': k, 'text': ' '.join(v[:2])} for k, v in points.items()]
    point_names = list(points)
    branches = []
    size = max(1, math.ceil(len(point_names) / 4))
    for i in range(0, len(point_names), size):
        group = point_names[i:i + size]
        branches.append({'title': f'{len(branches)+1:02d} ' + group[0], 'items': group})
    recap = paper.get('recap', [])
    return {'id': note_id or 'deeptutor-exam-' + digest, 'title': paper['title'],
            'question': '如何理解并运用《' + paper['title'] + '》中的知识点？',
            'source': paper['scope'], 'minutes': max(1, len(questions) * 2),
            'summary': summary, 'actions': [
                {'label': '主动回忆', 'text': '合上答案，复述这份试卷涉及的核心概念。'},
                {'label': '核对疑问', 'text': '对照本卷解析和原笔记核对错题与不确定之处。'},
                {'label': '再次练习', 'text': '给不确定的题目标记待复习，稍后重新作答。'}],
            'branches': branches, 'questions': questions, 'keywords': point_names[:3],
            'summaryLead': '；'.join(recap) or '围绕本卷知识点练习并核对参考答案。',
            'summaryBoundary': paper['scope'] + '。本页依据已保存试卷整理，知识总结来自题目解析，未重新读取原始资料；开放题请对照参考要点自评。'}


def export_html(paper, note_id=None):
    from pathlib import Path
    template = (Path(__file__).parent / 'assets' / 'exam-interactive-template.html').read_text(encoding='utf-8')
    payload = json.dumps(interactive_note(paper, note_id), ensure_ascii=False)
    # Prevent text from terminating the trusted inline script or starting HTML tags.
    payload = payload.replace('&', '\\u0026').replace('<', '\\u003c').replace('>', '\\u003e').replace('\u2028', '\\u2028').replace('\u2029', '\\u2029')
    return template.replace('__DEEPTUTOR_EXAM_DATA__', payload)
