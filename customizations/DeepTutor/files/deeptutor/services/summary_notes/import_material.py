"""Local material extraction; importing never invokes a paid LLM."""
import io
import json
import os
from pathlib import Path
import subprocess
import tempfile
import zipfile

SUPPORTED = {'.pdf', '.docx', '.pptx', '.xlsx', '.txt', '.md', '.csv', '.srt', '.vtt', '.mp3', '.wav', '.m4a', '.aac', '.flac', '.ogg', '.mp4', '.webm'}
IMAGES = {'.png', '.jpg', '.jpeg', '.webp', '.bmp', '.tif', '.tiff'}
SUPPORTED |= IMAGES
MEDIA = {'.mp3', '.wav', '.m4a', '.aac', '.flac', '.ogg', '.mp4', '.webm'}

def extract_material(name: str, data: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError('不支持此格式。请将旧版 DOC/PPT/XLS 转为 DOCX/PPTX/XLSX；图片支持PNG、JPEG、WebP、BMP、TIFF。')
    parts = []
    def add(value):
        parts.append(str(value))
        if sum(len(p) for p in parts) + len(parts) > 40000:
            raise ValueError('提取内容超过4万字，请拆分文件后导入；没有截断或发送模型。')
    if suffix in {'.docx', '.pptx', '.xlsx'}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 64 * 1024 * 1024:
                raise ValueError('Office 文件解压后过大，请拆分文件。')
    if suffix == '.pdf':
        import pymupdf
        from deeptutor.services.summary_notes.ocr import image_text
        with pymupdf.open(stream=data, filetype='pdf') as doc:
            if doc.needs_pass:
                raise ValueError('PDF 已加密，请先解锁后导入。')
            if len(doc) > 100:
                raise ValueError('PDF超过100页，请拆分后导入；没有跳过或截断页面。')
            document_has_text = False
            for index, page in enumerate(doc):
                text = page.get_text(sort=True).strip()
                add(f'【PDF第{index + 1}页】')
                if text:
                    document_has_text = True
                    add(text)
                rects = [page.rect] if not text else [
                    pymupdf.Rect(info['bbox']) & page.rect for info in page.get_image_info()
                ]
                seen = set()
                page_has_ocr = False
                for rect in rects:
                    key = tuple(round(v, 1) for v in rect)
                    if key in seen or rect.is_empty or rect.width < 15 or rect.height < 15:
                        continue
                    seen.add(key)
                    if len(seen) > 50:
                        raise ValueError('单页图片过多，请拆分或简化PDF后导入。')
                    # Crop long screenshot PDFs BEFORE rendering/resizing, preserving font size.
                    scale = min(6.0, max(2.0, 1200 / rect.width), 3600 / rect.width)
                    strip_height = 1800 / scale
                    import math
                    count = math.ceil(rect.height / strip_height)
                    if count > 100:
                        raise ValueError('单页截图过长，请拆成较短PDF后导入；没有截断。')
                    previous_lines = []
                    for strip_index in range(count):
                        top = rect.y0 + strip_index * strip_height
                        clip = pymupdf.Rect(rect.x0, top, rect.x1, min(rect.y1, top + strip_height + 12))
                        pix = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), clip=clip, alpha=False)
                        result = image_text(pix.pil_image())
                        if result:
                            page_has_ocr = True
                            document_has_text = True
                            lines = result.splitlines()
                            # Remove only matching adjacent overlap, never globally deduplicate text.
                            overlap = 0
                            for size in range(1, min(3, len(previous_lines), len(lines)) + 1):
                                if previous_lines[-size:] == lines[:size]: overlap = size
                            previous_lines = lines
                            if lines[overlap:]: add('【图片文字OCR，需核对】\n' + '\n'.join(lines[overlap:]))
                if not text and not page_has_ocr:
                    add('【本页未识别到文字，可能为空白、纯图或清晰度不足，请核对原页】')
            if not document_has_text:
                raise ValueError('PDF未识别到可用文字，请检查扫描清晰度或是否为纯图片内容。')
    elif suffix in IMAGES:
        from deeptutor.services.summary_notes.ocr import image_text
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            if getattr(image, 'n_frames', 1) > 100:
                raise ValueError('图片超过100帧，请拆分后导入。')
            for frame in range(getattr(image, 'n_frames', 1)):
                image.seek(frame)
                result = image_text(image)
                if result: add('【图片文字OCR，需核对】\n' + result)
    elif suffix == '.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        doc = Document(io.BytesIO(data))
        for block in doc.element.body:
            if block.tag.endswith('}p'): add(Paragraph(block, doc).text)
            elif block.tag.endswith('}tbl'):
                for row in Table(block, doc).rows: add(' | '.join(c.text for c in row.cells))
    elif suffix == '.pptx':
        from pptx import Presentation
        for slide in Presentation(io.BytesIO(data)).slides:
            for shape in slide.shapes:
                if shape.has_text_frame: add(shape.text)
                if shape.has_table:
                    for row in shape.table.rows: add(' | '.join(c.text for c in row.cells))
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        try:
            for sheet in book:
                add(sheet.title)
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index > 100000: raise ValueError('表格行数过多，请拆分文件。')
                    text = ' | '.join(str(c) if c is not None else '' for c in row)
                    if text.strip(' |'): add(text)
        finally: book.close()
    elif suffix in MEDIA:
        from deeptutor.services.path_service import get_path_service
        root_value = os.environ.get('SUMMARY_NOTES_ASR_ROOT', '')
        if not root_value:
            config = get_path_service().get_settings_file('summary_notes_asr')
            if config.is_file(): root_value = json.loads(config.read_text(encoding='utf-8')).get('root', '')
        if not root_value: raise ValueError('本机转写引擎未配置，请配置 SUMMARY_NOTES_ASR_ROOT 或 summary_notes_asr.json。')
        root = Path(root_value)
        python = root / 'data/crawler-python/Scripts/python.exe'
        if not python.is_file(): raise ValueError('本机转写引擎未配置，请先安装拾页抓取引擎或配置 SUMMARY_NOTES_ASR_ROOT。')
        with tempfile.TemporaryDirectory(prefix='note-import-') as temp:
            path = Path(temp) / ('material' + suffix)
            path.write_bytes(data)
            job = {'mode': 'transcribe', 'path': str(path), 'modelPath': str(root / 'data/crawler-models/base'), 'maxSeconds': 3600}
            result = subprocess.run([str(python), str(root / 'crawler/engine.py')], input=json.dumps(job), capture_output=True, text=True, encoding='utf-8', timeout=600, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
            try: payload = json.loads(result.stdout)
            except ValueError: raise ValueError('本机音视频转写失败，请检查文件和转写引擎。') from None
            if not payload.get('ok'): raise ValueError('音视频未能转写，请确认文件含有效声音，时长不超过60分钟。')
            add(payload['result']['content'])
    else:
        for encoding in ('utf-8-sig', 'utf-16', 'gb18030'):
            try:
                text = data.decode(encoding)
                if '\x00' in text: continue
                add(text)
                break
            except UnicodeError: continue
        else: raise ValueError('无法识别文字编码，请转换为 UTF-8 后导入。')
    text = '\n\n'.join(parts).strip()
    if len(text) > 40000: raise ValueError('提取内容超过4万字，请拆分文件后导入；没有截断或发送模型。')
    if not text: raise ValueError('未提取到可用文字。请检查图片或扫描页清晰度；音视频需包含清晰语音。')
    return text
